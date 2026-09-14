from pathlib import Path

from docx import Document


SOURCE = Path(
    r"D:\Documentary\xwechat_files\qizixiao_4688\msg\file\2026-09\AgeTogether_Iteration2_Bilingual_Speaker_Script.docx"
)
OUTPUT = Path(__file__).parent / "AgeTogether_Iteration2_Bilingual_Speaker_Script_Zixiao_edited.docx"


REPLACEMENTS = {
    135: (
        "Our AI design documents describe language and reading-style preferences, server-side task prompts, and visible failure behaviour. "
        "The flow is simple: a user taps a quick question or types their own question. "
        "The server adds the safety, language, and reading-style rules, then sends the request to DeepSeek. "
        "The browser shows a loading message, then an answer or a clear error. "
        "We tested an English quick reply and a Chinese typed reply on the deployed site. "
        "This is pretrained-model inference with prompt controls. "
        "We have not trained or fine-tuned a model, and we do not retrieve information from our place database. "
        "Separate provider timeout handling, consent wording, and broader safety testing are still next steps."
    ),
    137: (
        "我们的 AI 设计文档说明了语言、阅读风格、服务端任务提示和可见的失败处理。"
        "流程很简单：用户点击快捷问题，或输入自己的问题。"
        "服务器加入安全、语言和阅读风格规则，再把请求发送给 DeepSeek。"
        "浏览器先显示等待提示，然后显示答案或清楚的错误信息。"
        "我们在部署网站上测试了英文快捷回复和中文输入回复。"
        "这是通过提示控制使用预训练模型进行推理；我们没有训练或微调模型，也没有从地点数据库检索信息。"
        "提供方超时处理、同意说明和更广泛的安全测试，仍是下一步工作。"
    ),
    143: (
        "Our innovation combines a familiar companion with language choices, simple reading styles, and local personalisation. "
        "A chosen photo can become the companion through background removal in the browser, with a simple crop as a fallback. "
        "The latest code adds short text bubbles for AI answers, a wellbeing tip after every ten Pet clicks, and configurable daily reminders. "
        "The bubble shows text, not spoken audio. "
        "Reminders are stored locally and work while the page is open; the prototype currently enables the default reminders. "
        "For a real release, we would use clear opt-in, especially for medication reminders. "
        "We will test whether these features help people complete a task or cause confusion. "
        "We will measure understanding and effort, rather than claim that the Pet reduces loneliness."
    ),
    145: (
        "我们的创新把熟悉的助手形象、语言选择、简单的阅读风格和本地个性化结合起来。"
        "用户选择的照片可以在浏览器中去除背景，失败时则使用简单裁切作为回退。"
        "最新代码为 AI 答案加入简短文字气泡，每点击 Pet 十次显示一条健康小贴士，并加入可配置的每日提醒。"
        "气泡是文字，不是语音。"
        "提醒保存在浏览器本地，只在页面打开时运行；目前原型默认开启这些提醒。"
        "正式发布时，我们会采用明确的用户选择，尤其是用药提醒。"
        "我们会测试这些功能是帮助用户完成任务，还是造成困惑。"
        "我们会衡量理解程度和操作负担，而不是宣称 Pet 能降低孤独感。"
    ),
    151: (
        "Our next steps are to close the confirmed navigation and location issues, complete the AI safety and failure checks, "
        "and test the highest-priority accessibility improvements. "
        "We will also test whether the Pet bubbles, health tips, and reminders are easy to understand and useful. "
        "Then we will review the results against the Must Have stories and our Definition of Done. "
        "Data freshness and real account permissions are still dependencies for a wider release. "
        "That brings us to the end of our Iteration 2 Product Discovery Presentation. Thank you."
    ),
    153: (
        "下一步，我们会先解决已确认的导航和位置问题，完成 AI 安全与失败处理检查，再测试最优先的无障碍改进。"
        "我们也会测试 Pet 气泡、健康小贴士和提醒是否容易理解、是否真正有帮助。"
        "随后，我们会根据 Must Have 用户故事和完成定义评审结果。"
        "扩大正式发布之前，数据时效性和真实账户权限仍是必要条件。"
        "以上就是我们 Iteration 2 产品发现展示的全部内容。感谢大家。"
    ),
    173: "Did you train or fine-tune the AI model?",
    174: (
        "We use pretrained-model inference through a server API and task prompts. "
        "We have not trained or fine-tuned a model, and the current flow does not retrieve from the place database."
    ),
    175: "我们通过服务器 API 和任务提示调用预训练模型，没有自行训练或微调，目前也没有检索地点数据库。",
}


def replace_paragraph(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


document = Document(SOURCE)
for index, replacement in REPLACEMENTS.items():
    replace_paragraph(document.paragraphs[index], replacement)
document.save(OUTPUT)
print(OUTPUT)
